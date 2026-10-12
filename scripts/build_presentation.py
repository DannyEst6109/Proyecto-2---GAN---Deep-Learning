"""Genera presentacion.pptx a partir de docs/presentacion_guion.md."""

import argparse
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BACKGROUND = RGBColor(0x16, 0x1A, 0x24)
PANEL = RGBColor(0x22, 0x28, 0x36)
ACCENT = RGBColor(0xE8, 0x6A, 0x33)
TEXT = RGBColor(0xEC, 0xEC, 0xEC)
MUTED = RGBColor(0x9A, 0xA3, 0xB5)
FONT = 'Segoe UI'


def read_slides(script):
    """'## título' abre una diapositiva; '- ' viñeta, '| a | b |' fila de tabla, '![](ruta)' imagen."""
    slides = []
    for line in script.read_text(encoding='utf-8').splitlines():
        if line.startswith('## '):
            slides.append({'title': line[3:], 'bullets': [], 'rows': [], 'image': None})
        elif slides and line.startswith('- '):
            slides[-1]['bullets'].append(line[2:])
        elif slides and line.startswith('|'):
            slides[-1]['rows'].append([cell.strip() for cell in line.strip('|').split('|')])
        elif slides and (match := re.match(r'!\[.*\]\((.+)\)', line)):
            slides[-1]['image'] = script.parent / match.group(1)
    return slides


def add_text(slide, text, left, top, width, height, size, color=TEXT, bold=False):
    frame = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    frame.word_wrap = True
    write(frame.paragraphs[0], text, size, color, bold)
    return frame


def write(paragraph, text, size, color=TEXT, bold=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
    run.font.color.rgb = color


def add_picture(slide, path, left, top, width, height):
    """Centra la imagen dentro del recuadro sin deformarla."""
    with Image.open(path) as picture:
        ratio = picture.width / picture.height
    if ratio > width / height:
        size = (width, width / ratio)
    else:
        size = (height * ratio, height)
    slide.shapes.add_picture(str(path), Inches(left + (width - size[0]) / 2), Inches(top + (height - size[1]) / 2),
                             Inches(size[0]), Inches(size[1]))


def add_table(slide, rows):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.3))
    for width, column in zip((3.3, 0.9, 4.3, 3.6), shape.table.columns):
        column.width = Inches(width)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = shape.table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else PANEL if i % 2 else BACKGROUND
            write(cell.text_frame.paragraphs[0], text, 12, TEXT, bold=i == 0)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--script', type=Path, default=Path('docs/presentacion_guion.md'))
    parser.add_argument('--output', type=Path, default=Path('presentacion.pptx'))
    args = parser.parse_args()

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
    slides = read_slides(args.script)
    for number, content in enumerate(slides, start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])  # En blanco.
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BACKGROUND
        add_text(slide, content['title'], 0.6, 0.35, 12.1, 0.8, 34 if number == 1 else 28, bold=True)
        bar = slide.shapes.add_shape(1, Inches(0.6), Inches(1.2), Inches(1.4), Inches(0.07))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        add_text(slide, f'Dragones de fantasía medieval · GAN   {number}', 8.4, 7.0, 4.4, 0.4, 10, MUTED)
        if content['rows']:
            add_table(slide, content['rows'])
            continue

        image = content['image'] if content['image'] and content['image'].exists() else None
        wide = False
        if image:
            with Image.open(image) as picture:
                wide = picture.width > 2 * picture.height
        # Imagen ancha: abajo a todo el ancho. Imagen normal: a la derecha.
        width = 12.1 if wide or not image else 6.0
        frame = add_text(slide, '', 0.6, 1.5, width, 2.6 if wide else 5.3, 16)
        for index, bullet in enumerate(content['bullets']):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_after = Pt(7)
            write(paragraph, '▸ ', 15 if image else 18, ACCENT, bold=True)
            write(paragraph, bullet, 15 if image else 18)
        if wide:
            add_picture(slide, image, 0.6, 4.0, 12.1, 2.9)
        elif image:
            add_picture(slide, image, 7.0, 1.5, 5.7, 5.3)
    deck.save(args.output)

    main_slides = sum(not content['title'].startswith('Anexo') for content in slides)
    print(f'{args.output}: {main_slides} diapositivas y {len(slides) - main_slides} anexos.')
    if main_slides > 12:
        print('AVISO: el máximo es 12 diapositivas sin contar anexos.')


if __name__ == '__main__':
    main()
